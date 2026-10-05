"""Genera la plantilla de Google Sheets (xlsx) de la réplica de niveles de procesamiento.
Uso:  python3 crear_plantilla_xlsx.py [salida.xlsx] [--demo]
  --demo  rellena 25 filas con los datos de un grupo de clase (solo para comprobar las fórmulas).
Las fórmulas usan solo SUMPRODUCT, SQRT, LN, FDIST y CHIDIST: funcionan igual en Google Sheets y en Excel."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
H = json.load(open(os.path.join(AQUI, "columnas.json")))
out = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else "Niveles_Procesamiento_datos.xlsx"
demo = "--demo" in sys.argv
AZUL, MARINO, NARANJA, GRIS = "0047E8", "001447", "FF9606", "43506B"
F = "Arial"
wb = Workbook()


def cab(ws, fila, textos, color=AZUL, col0=1):
    for i, t in enumerate(textos, col0):
        c = ws.cell(fila, i, t)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)


# ---------------- Respuestas ----------------
ws = wb.active; ws.title = "Respuestas"
cab(ws, 1, H); ws.freeze_panes = "B2"
for i, h in enumerate(H, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(9, min(24, len(h) + 3))
col = {h: get_column_letter(i + 1) for i, h in enumerate(H)}
N = 2000
def R(h): return f"Respuestas!${col[h]}$2:${col[h]}${N}"

if demo:
    # Datos de un grupo de clase (N = 25): los mismos que se analizaron en JASP
    X = [[.5, .66, .83], [.33, .33, .5], [.5, .5, .83], [.83, .66, .66], [.66, .66, .83], [.33, .16, .66], [.5, .17, .5],
         [.66, .66, .66], [.7, 0, .5], [.17, .5, 1], [.66, .33, 1], [.5, .5, 1], [.33, .83, 1], [.5, .66, .66], [.33, .5, .66],
         [.5, .5, .83], [.16, .3, .83], [.66, .16, .66], [.5, .83, 1], [.83, .66, .83], [0, .5, .66], [.5, .667, .833],
         [.33, .66, 1], [.83, .66, .5], [.3, .33, 1]]
    for r, x in enumerate(X, 2):
        fila = {h: "" for h in H}
        fila.update(id=f"DEMO-{r - 1}", cohorte="MDTI-2026-G1", sexo="Hombre" if r % 3 == 0 else "Mujer", edad=20,
                    superficial=x[0], intermedio=x[1], profundo=x[2], sospecha="No", conocia="No", condiciones="Sí",
                    cambios_pestana=0, falsas_alarmas=(r * 7) % 5)
        for i, h in enumerate(H, 1): ws.cell(r, i, fila[h])

# ---------------- Resumen (fórmulas en vivo) ----------------
rs = wb.create_sheet("Resumen", 0)
rs.column_dimensions["A"].width = 44
for c in "BCDEFG": rs.column_dimensions[c].width = 16
rs["A1"] = "Niveles de procesamiento y recuerdo · Resumen en vivo"
rs["A1"].font = Font(name=F, size=16, bold=True, color=AZUL)
rs["A2"] = "ESIC University · Métodos, Diseño y Técnicas de Investigación en Psicología · Tema 5 (diseño con el mismo grupo)"
rs["A2"].font = Font(name=F, size=10, color=MARINO)
rs["A4"] = "Cohorte a analizar (vacío = todas):"; rs["B4"] = ""
rs["B4"].fill = PatternFill("solid", fgColor="FFF3E0"); rs["B4"].font = Font(name=F, bold=True)
rs["A5"] = ("Excluidos automáticamente: datos incompletos, sospechaban la prueba de memoria, conocían el experimento "
            "o tuvieron interrupciones. El análisis completo (Mauchly, post hoc, Friedman) está en la pestaña «Resultados» de la web.")
rs["A5"].font = Font(name=F, size=9, italic=True, color=GRIS); rs["A5"].alignment = Alignment(wrap_text=True)
rs.merge_cells("A5:G5"); rs.row_dimensions[5].height = 30

S, I, P = R("superficial"), R("intermedio"), R("profundo")
MASK = (f'({S}<>"")*({I}<>"")*({P}<>"")*({R("sospecha")}<>"Sí")*({R("conocia")}<>"Sí")*({R("condiciones")}<>"No")'
        f'*((({R("cohorte")}=$B$4)+($B$4=""))>0)')

# Descriptivos
cab(rs, 7, ["Descriptivos (proporción reconocida)", "Superficial", "Intermedio", "Profundo"])
lv = {"B": S, "C": I, "D": P}
rs["A8"] = "n incluidos"
for c in "BCD": rs[f"{c}8"] = f"=SUMPRODUCT({MASK})"
rs["A9"] = "Media"; rs["A10"] = "Desviación típica"; rs["A11"] = "Error típico"; rs["A12"] = "Coeficiente de variación"
for c, x in lv.items():
    rs[f"{c}9"] = f'=IFERROR(SUMPRODUCT({MASK}*{x})/{c}8,"")'
    rs[f"{c}10"] = f'=IFERROR(SQRT(SUMPRODUCT({MASK}*({x}-{c}9)^2)/({c}8-1)),"")'
    rs[f"{c}11"] = f'=IFERROR({c}10/SQRT({c}8),"")'
    rs[f"{c}12"] = f'=IFERROR({c}10/{c}9,"")'
for r in range(9, 13):
    for c in "BCD": rs[f"{c}{r}"].number_format = "0.000"

# Por sexo
cab(rs, 14, ["Media por sexo", "Superficial", "Intermedio", "Profundo", "n"])
for k, sx in enumerate(["Mujer", "Hombre"]):
    r = 15 + k; m = f'{MASK}*({R("sexo")}="{sx}")'
    rs[f"A{r}"] = sx
    rs[f"E{r}"] = f"=SUMPRODUCT({m})"
    for c, x in lv.items():
        rs[f"{c}{r}"] = f'=IFERROR(SUMPRODUCT({m}*{x})/$E{r},"")'; rs[f"{c}{r}"].number_format = "0.000"

# ANOVA de medidas repetidas (esfericidad asumida + Greenhouse-Geisser)
cab(rs, 18, ["ANOVA de medidas repetidas", "SC", "gl", "MC", "F", "p", ""])
n = "$B$8"; g = "(($B$9+$C$9+$D$9)/3)"
sstot = f"SUMPRODUCT({MASK}*(({S}-{g})^2+({I}-{g})^2+({P}-{g})^2))"
sssuj = f"3*SUMPRODUCT({MASK}*((({S}+{I}+{P})/3-{g})^2))"
rs["A19"] = "Nivel de procesamiento (esfericidad asumida)"
rs["B19"] = f'=IFERROR({n}*(($B$9-{g})^2+($C$9-{g})^2+($D$9-{g})^2),"")'
rs["C19"] = 2; rs["D19"] = '=IFERROR(B19/C19,"")'; rs["E19"] = '=IFERROR(D19/D20,"")'
rs["F19"] = '=IFERROR(FDIST(E19,C19,C20),"")'
rs["A20"] = "Residuos (error)"
rs["B20"] = f'=IFERROR({sstot}-B19-B21,"")'; rs["C20"] = f'=IFERROR(2*({n}-1),"")'; rs["D20"] = '=IFERROR(B20/C20,"")'
rs["A21"] = "Entre sujetos (diferencias individuales)"
rs["B21"] = f'=IFERROR({sssuj},"")'; rs["C21"] = f'=IFERROR({n}-1,"")'; rs["D21"] = '=IFERROR(B21/C21,"")'
rs["A22"] = "ω² (como JASP)"; rs["B22"] = '=IFERROR(MAX(0,C19*(D19-D20)/(B19+B20+B21+D21)),"")'
rs["A23"] = "η² parcial"; rs["B23"] = '=IFERROR(B19/(B19+B20),"")'
for r in range(19, 24):
    for c in "BDEF": rs[f"{c}{r}"].number_format = "0.000"

# Esfericidad: matriz de covarianzas y contrastes de Helmert
cab(rs, 25, ["Esfericidad (Mauchly) y ε", "Valor", "", "", "", "", ""])
cov = {}
for a, xa in lv.items():
    for b, xb in lv.items():
        cov[(a, b)] = f"(SUMPRODUCT({MASK}*({xa}-${a}$9)*({xb}-${b}$9))/({n}-1))"
s11, s22, s33 = cov[("B", "B")], cov[("C", "C")], cov[("D", "D")]
s12, s13, s23 = cov[("B", "C")], cov[("B", "D")], cov[("C", "D")]
# h1 = (1,-1,0)/√2 ; h2 = (1,1,-2)/√6
rs["A26"] = "Sc11"; rs["B26"] = f'=IFERROR(({s11}+{s22}-2*{s12})/2,"")'
rs["A27"] = "Sc22"; rs["B27"] = f'=IFERROR(({s11}+{s22}+4*{s33}+2*{s12}-4*{s13}-4*{s23})/6,"")'
rs["A28"] = "Sc12"; rs["B28"] = f'=IFERROR(({s11}-{s22}-2*{s13}+2*{s23})/SQRT(12),"")'
rs["A29"] = "W de Mauchly"; rs["B29"] = '=IFERROR((B26*B27-B28^2)/((B26+B27)/2)^2,"")'
rs["A30"] = "χ² aproximado (gl = 2)"; rs["B30"] = f'=IFERROR(-({n}-2)*LN(B29),"")'
rs["A31"] = "p"; rs["B31"] = '=IFERROR(CHIDIST(B30,2),"")'
rs["A32"] = "ε de Greenhouse-Geisser"; rs["B32"] = '=IFERROR(MIN(1,(B26+B27)^2/(2*(B26^2+2*B28^2+B27^2))),"")'
rs["A33"] = "gl corregidos (efecto · error)"; rs["B33"] = '=IFERROR(C19*B32,"")'; rs["C33"] = '=IFERROR(C20*B32,"")'
rs["A34"] = "p con Greenhouse-Geisser (gl redondeados)"; rs["B34"] = '=IFERROR(FDIST(E19,ROUND(B33,0),ROUND(C33,0)),"")'
for r in range(26, 35):
    rs[f"B{r}"].number_format = "0.000"
rs["C33"].number_format = "0.000"
rs["A36"] = ("La p exacta con gl decimales (Greenhouse-Geisser, Huynh-Feldt), las comparaciones post hoc y la prueba de Friedman "
             "se calculan en la pestaña «Resultados» de la web.")
rs["A36"].font = Font(name=F, size=9, italic=True, color=GRIS); rs.merge_cells("A36:G36"); rs["A36"].alignment = Alignment(wrap_text=True)
rs.row_dimensions[36].height = 28
rs["A37"] = "Referencia (grupo de clase, JASP): M = 0.48 / 0.50 / 0.78; F(1.74, 41.87) = 17.89, p < .001, ω² = .29."
rs["A37"].font = Font(name=F, size=9, italic=True, color=GRIS)
for row in rs.iter_rows(min_row=8, max_row=34):
    for c in row:
        if c.font.color is None or c.font.color.rgb != "00FFFFFF":
            c.font = Font(name=F, bold=c.font.bold)

# ---------------- Estímulos ----------------
cfg_txt = open(os.path.join(AQUI, "..", "js", "config.js"), encoding="utf-8").read()
import re
pregs = re.findall(r'id: "(q\d\d)", nivel: "(\w+)",\s+palabra: "([^"]+)",\s+texto: "([^"]+)",\s+correcta: "([^"]+)"', cfg_txt)
es = wb.create_sheet("Estimulos")
cab(es, 1, ["Orden", "Pregunta", "Nivel de la VI", "Palabra", "Respuesta correcta", "Bloque"])
for w, c in zip([8, 62, 16, 12, 18, 9], "ABCDEF"): es.column_dimensions[c].width = w
for i, (qid, nivel, pal, txt, ok) in enumerate(pregs, 2):
    for j, v in enumerate([i - 1, txt, nivel, pal, ok, (i - 2) // 3 + 1], 1): es.cell(i, j, v)
lista = re.search(r"RECONOCIMIENTO: \[(.*?)\]", cfg_txt, re.S).group(1)
palabras = re.findall(r'"([^"]+)"', lista)
r0 = len(pregs) + 4
cab(es, r0, ["Lista de reconocimiento (36 palabras: 18 vistas + 18 nuevas)"] + [""] * 5, MARINO)
for k, pw in enumerate(palabras):
    es.cell(r0 + 1 + k // 6, 1 + k % 6, pw)

# ---------------- Leeme ----------------
lm = wb.create_sheet("Leeme"); lm.column_dimensions["A"].width = 30; lm.column_dimensions["B"].width = 95
cab(lm, 1, ["Pestaña / columna", "Descripción"], MARINO)
texto = [
    ("Resumen", "Estadísticos en vivo calculados con fórmulas a partir de «Respuestas». Escribe una cohorte en B4 para filtrar."),
    ("Respuestas", "Una fila por participante. La escribe la web: no cambies las cabeceras de la fila 1."),
    ("Estimulos", "Las 18 preguntas (orden, nivel, palabra y respuesta correcta) y la lista de reconocimiento."),
    ("superficial · intermedio · profundo", "VD: proporción de palabras reconocidas en cada nivel (aciertos / 6). Formato ancho: una columna por nivel, como en JASP."),
    ("aciertos_<nivel>", "Número de palabras reconocidas en ese nivel (0-6)."),
    ("falsas_alarmas · p_falsas_alarmas", "Palabras nuevas marcadas como reconocidas (0-18) y su proporción."),
    ("qXX_resp · qXX_ok · qXX_tr_ms · qXX_rec", "Tarea 1: respuesta Sí/No, acierto en la pregunta, tiempo de respuesta (ms) y si la palabra se reconoció después (1/0)."),
    ("orden_preguntas", "Orden en que se presentaron las 18 preguntas (aleatorización por bloques)."),
    ("t2_*", "Tarea distractora: números correctos, errores, si la completó y tiempo empleado."),
    ("sexo · edad", "Variables descriptivas (permiten separar los descriptivos por sexo, como en JASP)."),
    ("sospecha · conocia · condiciones · cambios_pestana", "Criterios de exclusión."),
    ("Apps Script", "Extensiones > Apps Script: pegar apps-script/Code.gs e implementar como aplicación web (ver README)."),
]
for i, (a, b) in enumerate(texto, 2):
    lm.cell(i, 1, a).font = Font(name=F, bold=True); lm.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
wb.save(out); print("ok", out)
